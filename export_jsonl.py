"""Export shahnameh_style_dataset.xlsx to chat-format JSONL for supervised fine-tuning.

Usage:
    python export_jsonl.py                      # all languages, with system prompt
    python export_jsonl.py --no-system          # omit the system message
    python export_jsonl.py --langs fa en        # only some languages

Writes <out>/<lang>_train.jsonl, <lang>_val.jsonl and all_train.jsonl, all_val.jsonl.
Every row is validated first; nothing is written if any row fails.
"""
import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent
SHEETS = {"fa": "Persian (FA)", "en": "English (EN)", "de": "German (DE)"}
SYSTEM_SHEET = "System Prompts"
REQUIRED_COLUMNS = ("id", "category", "split", "user_prompt", "assistant_response")
SPLITS = ("train", "val")
ID_PATTERN = re.compile(r"^SH-\d{3,}$")
PERSIAN_SCRIPT = re.compile(r"[؀-ۿ]")


def read_sheet(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        raise SystemExit(f"Sheet '{sheet_name}' not found in workbook")
    rows = wb[sheet_name].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    return header, [
        (n, {h: "" if v is None else str(v).strip() for h, v in zip(header, values)})
        for n, values in enumerate(rows, start=2)
        if any(v not in (None, "") for v in values)
    ]


def load_system_prompts(wb):
    _, rows = read_sheet(wb, SYSTEM_SHEET)
    return {r.get("lang", "").lower(): r.get("system_prompt", "") for _, r in rows}


def validate_rows(lang, header, rows):
    sheet = SHEETS[lang]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        return [f"{sheet}: missing columns {missing}"], []
    errors, warnings, seen = [], [], set()
    for n, row in rows:
        where = f"{sheet} row {n}"
        if not ID_PATTERN.match(row["id"]):
            errors.append(f"{where}: id {row['id']!r} should look like SH-001")
        elif row["id"] in seen:
            errors.append(f"{where}: duplicate id {row['id']}")
        seen.add(row["id"])
        if row["split"] not in SPLITS:
            errors.append(f"{where}: split must be one of {SPLITS}, got {row['split']!r}")
        errors.extend(f"{where}: {col} is empty" for col in ("user_prompt", "assistant_response") if not row[col])
        text = row["user_prompt"] + row["assistant_response"]
        if lang == "fa" and row["assistant_response"] and not PERSIAN_SCRIPT.search(row["assistant_response"]):
            errors.append(f"{where}: Persian response contains no Persian script")
        if lang != "fa" and PERSIAN_SCRIPT.search(text):
            errors.append(f"{where}: {lang.upper()} row contains Persian script")
        if row["assistant_response"] and len(row["assistant_response"].splitlines()) % 2:
            warnings.append(f"{where}: odd number of verse lines (couplets expected)")
    return errors, warnings


def to_example(row, system_prompt):
    system = [{"role": "system", "content": system_prompt}] if system_prompt else []
    return {"messages": system + [
        {"role": "user", "content": row["user_prompt"]},
        {"role": "assistant", "content": row["assistant_response"]},
    ]}


def write_jsonl(path, examples):
    with path.open("w", encoding="utf-8", newline="\n") as f:
        f.writelines(json.dumps(ex, ensure_ascii=False) + "\n" for ex in examples)


def parse_args(argv):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", type=Path, default=HERE / "shahnameh_style_dataset.xlsx")
    parser.add_argument("--out", type=Path, default=HERE / "training_data")
    parser.add_argument("--langs", nargs="+", choices=sorted(SHEETS), default=list(SHEETS))
    parser.add_argument("--no-system", action="store_true", help="leave out the system message")
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    if not args.xlsx.exists():
        raise SystemExit(f"Workbook not found: {args.xlsx}")
    wb = load_workbook(args.xlsx, read_only=True)
    prompts = {} if args.no_system else load_system_prompts(wb)

    sheets, errors, warnings = {}, [], []
    for lang in args.langs:
        header, rows = read_sheet(wb, SHEETS[lang])
        sheet_errors, sheet_warnings = validate_rows(lang, header, rows)
        errors += sheet_errors
        warnings += sheet_warnings
        sheets[lang] = [row for _, row in rows]
    wb.close()

    id_sets = {lang: {r["id"] for r in rows} for lang, rows in sheets.items()}
    if len({frozenset(ids) for ids in id_sets.values()}) > 1:
        warnings.append("id sets differ between language sheets (rows are no longer parallel)")
    for message in warnings:
        print(f"warning: {message}")
    if errors:
        print("\n".join(f"error: {e}" for e in errors), file=sys.stderr)
        print(f"{len(errors)} error(s); nothing written.", file=sys.stderr)
        return 1

    args.out.mkdir(parents=True, exist_ok=True)
    combined = {split: [] for split in SPLITS}
    for lang, rows in sheets.items():
        for split in SPLITS:
            tagged = [(r["id"], lang, to_example(r, prompts.get(lang, ""))) for r in rows if r["split"] == split]
            write_jsonl(args.out / f"{lang}_{split}.jsonl", [ex for _, _, ex in tagged])
            combined[split] += tagged
            print(f"{lang}_{split}.jsonl: {len(tagged)} examples")
    for split, tagged in combined.items():
        write_jsonl(args.out / f"all_{split}.jsonl", [ex for _, _, ex in sorted(tagged, key=lambda t: t[:2])])
        print(f"all_{split}.jsonl: {len(tagged)} examples")
    return 0


if __name__ == "__main__":
    sys.exit(main())
